#!/usr/bin/env python3
"""Сборка PPTX-версии презентации из отрендеренного PDF.

Каждая страница PDF переносится на слайд формата 16:9 в разрешении 1920x1080.
Такая версия предназначена для показа в PowerPoint или Google Slides;
для правки содержания используется index.html с последующей пересборкой.
"""

import pathlib
import tempfile

import pymupdf
from pptx import Presentation
from pptx.util import Inches

DIR = pathlib.Path(__file__).resolve().parent
PDF = DIR / "Rail-Logistics-Hub-Astana.pdf"
OUT = DIR / "Rail-Logistics-Hub-Astana.pptx"
DPI = 144  # 1280x720 CSS-пикселей -> 1920x1080


def main() -> None:
    if not PDF.exists():
        raise SystemExit(f"Сначала соберите PDF: {PDF}")

    doc = pymupdf.open(PDF)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory() as tmp:
        for i, page in enumerate(doc):
            png = pathlib.Path(tmp) / f"slide_{i + 1:02d}.png"
            page.get_pixmap(dpi=DPI).save(png)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(png), 0, 0, width=prs.slide_width, height=prs.slide_height
            )

    prs.save(OUT)
    print(f"PPTX готов: {OUT} ({len(doc)} слайдов)")


if __name__ == "__main__":
    main()
