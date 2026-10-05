#!/usr/bin/env python3
"""PPTX для показа: каждый слайд PDF становится картинкой 16:9."""

import pathlib
import tempfile

import pymupdf
from pptx import Presentation
from pptx.util import Inches

DIR = pathlib.Path(__file__).resolve().parent
PDF = DIR / "Wagon-Park-Kazakhstan.pdf"
OUT = DIR / "Wagon-Park-Kazakhstan.pptx"


def main() -> None:
    doc = pymupdf.open(PDF)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    with tempfile.TemporaryDirectory() as tmp:
        for i, page in enumerate(doc):
            png = pathlib.Path(tmp) / f"s{i+1:02d}.png"
            page.get_pixmap(dpi=144).save(png)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(OUT)
    print(f"PPTX готов: {OUT} ({len(doc)} слайдов)")


if __name__ == "__main__":
    main()
